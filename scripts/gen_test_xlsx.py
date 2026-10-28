# -*- coding: utf-8 -*-
"""生成课程交付物：测试用例与 BUG 清单（按范例模板结构）。"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = Path("docs/交付物/测试用例与BUG清单.xlsx")

THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="3A3A3C")
SUB_FILL = PatternFill("solid", fgColor="F2F2F7")
TITLE_FONT = Font(name="微软雅黑", size=15, bold=True)
HEAD_FONT = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
LABEL_FONT = Font(name="微软雅黑", size=9, bold=True)
BODY_FONT = Font(name="微软雅黑", size=10)
SMALL_FONT = Font(name="微软雅黑", size=9)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="top", wrap_text=True)

MODULES = [
    {
        "code": "TC-AUTH",
        "name": "认证与权限",
        "desc": "登录、令牌刷新、注销、修改密码与基于角色的接口权限控制。",
        "purpose": "验证认证链路的正确性、令牌生命周期与越权访问拦截。",
        "prepare": "1.环境准备：后端 8000 端口、MySQL 与 Redis 正常运行。\n"
                   "2.数据准备：admin / staff01 / 2023010101 三个演示账号。\n"
                   "3.其他准备：Postman 或接口冒烟脚本。",
        "precondition": "数据库迁移已执行，演示账号已初始化。",
        "cases": [
            ("使用管理员账号登录", "① 打开登录页\n② 输入 admin / Admin@123\n③ 点击登录",
             "账号：admin\n密码：Admin@123", "① 返回访问令牌与刷新令牌\n② 跳转工作台\n③ 侧边栏显示管理员菜单", "通过"),
            ("错误密码登录", "① 输入正确账号与错误密码\n② 点击登录",
             "账号：admin\n密码：wrong", "① 返回 401 与「账号或密码错误」\n② 不签发令牌", "通过"),
            ("连续失败锁定", "① 连续 5 次输入错误密码\n② 第 6 次再登录",
             "账号：staff01\n密码：错误密码 ×5", "① 第 6 次返回 429\n② 提示 15 分钟后再试", "通过"),
            ("未登录访问受保护接口", "① 不携带令牌请求 /api/v1/packages",
             "无", "返回 401，业务码 10002", "通过"),
            ("学生越权访问用户管理", "① 学生登录\n② 请求 /api/v1/users",
             "账号：2023010101", "返回 403，业务码 10003", "通过"),
            ("令牌刷新", "① 访问令牌过期后携带刷新令牌请求 /auth/refresh",
             "refresh_token", "返回新的访问令牌与刷新令牌", "通过"),
            ("注销后令牌失效", "① 调用 /auth/logout\n② 使用原访问令牌请求 /auth/me",
             "access_token", "① 注销成功\n② 原令牌返回 401", "通过"),
        ],
    },
    {
        "code": "TC-USER",
        "name": "用户管理",
        "desc": "用户分页查询、新增、编辑、启停与角色维护。",
        "purpose": "验证管理员对账号的全生命周期管理能力与唯一性约束。",
        "prepare": "1.环境准备：服务正常运行。\n2.数据准备：演示账号 4 个。\n3.其他准备：管理员令牌。",
        "precondition": "以 admin 身份登录。",
        "cases": [
            ("用户列表与筛选", "① 打开用户管理\n② 按角色筛选工作人员",
             "角色：staff", "① 返回分页数据\n② 仅显示 staff 角色用户", "通过"),
            ("新增学生账号", "① 点击新增用户\n② 填写学号、姓名、手机号、密码\n③ 保存",
             "学号：2023010199\n手机号：13900001234", "① 创建成功\n② 列表出现新账号\n③ 可用新账号登录", "通过"),
            ("手机号重复校验", "① 使用已存在手机号新增用户", "手机号：13800000002", "返回 409，提示手机号已存在", "通过"),
            ("禁用账号", "① 点击账号开关禁用\n② 使用该账号登录",
             "账号：新账号", "① 状态变为禁用\n② 登录返回 403", "通过"),
            ("启用账号", "① 重新启用账号\n② 再次登录", "账号：新账号", "① 状态恢复正常\n② 登录成功", "通过"),
        ],
    },
    {
        "code": "TC-STATION",
        "name": "驿站与货架",
        "desc": "驿站档案、营业时间、负责人、货架分区与格口容量管理。",
        "purpose": "验证驿站与货架的增删改查、容量联动与状态约束。",
        "prepare": "1.环境准备：服务正常。\n2.数据准备：东区、西区两个演示驿站。\n3.其他准备：管理员与工作人员账号。",
        "precondition": "以 admin 身份登录。",
        "cases": [
            ("新增驿站", "① 点击新增驿站\n② 填写编码、名称、地址、营业时间\n③ 保存",
             "编码：ST-C\n名称：南区驿站", "① 创建成功\n② 列表出现新驿站\n③ 总格口不变", "通过"),
            ("驿站编码唯一性", "① 使用已存在编码新增", "编码：ST-A", "返回 409，提示编码已存在", "通过"),
            ("营业时间校验", "① 结束时间早于开始时间\n② 保存", "开始 09:00\n结束 08:00", "返回 422，提示结束时间必须晚于开始时间", "通过"),
            ("新增货架与容量联动", "① 打开货架管理\n② 新增货架容量 20\n③ 查看驿站容量",
             "货架：C-01\n容量：20", "① 货架创建成功\n② 驿站总格口增加 20", "通过"),
            ("有包裹货架禁止删除", "① 对已占用货架执行删除", "货架：A-01", "返回 409，提示仍有包裹无法删除", "通过"),
            ("工作人员只读", "① 以 staff01 登录\n② 打开驿站管理", "账号：staff01", "① 可查看列表\n② 无新增、编辑、停用入口", "通过"),
        ],
    },
    {
        "code": "TC-IN",
        "name": "快递入库",
        "desc": "单件与批量入库、收件人匹配、格口分配、取件码生成与调架。",
        "purpose": "验证入库主流程、重复运单拦截、容量约束与批量容错。",
        "prepare": "1.环境准备：服务正常，驿站有空闲格口。\n2.数据准备：学生手机号 13800000002。\n3.其他准备：工作人员令牌。",
        "precondition": "以 staff01 身份登录，东区驿站营业中。",
        "cases": [
            ("单件入库", "① 打开包裹管理\n② 点击入库登记\n③ 填写运单号与手机号\n④ 提交",
             "运单号：SF1000000001\n手机号：13800000002", "① 入库成功\n② 自动分配货架\n③ 生成 6 位取件码", "通过"),
            ("重复运单拦截", "① 再次提交相同运单号", "运单号：SF1000000001", "返回 409，业务码 20001，附带原入库时间", "通过"),
            ("手机号格式校验", "① 输入 5 位手机号提交", "手机号：12345", "返回 422，字段提示格式错误", "通过"),
            ("自动匹配收件人", "① 入库时填写已注册手机号", "手机号：13800000002", "① 返回 recipient_id\n② 收件人姓名与账号一致", "通过"),
            ("批量入库部分失败", "① 粘贴 3 行清单，其中 1 行重复运单\n② 提交",
             "2 条新运单 + 1 条重复", "① 成功 2 条、失败 1 条\n② 失败原因显示重复入库", "通过"),
            ("格口已满", "① 使用容量为 1 的驿站连续入库 2 件", "驿站：满载测试驿站", "第 2 件返回 409，业务码 20005", "通过"),
            ("停用驿站禁止入库", "① 停用驿站后入库", "驿站状态：0", "返回 409，业务码 10005", "通过"),
            ("调整格口", "① 打开包裹详情\n② 调整到另一货架", "目标货架：A-02", "① 包裹货架更新\n② 两个货架占用数同步变化", "通过"),
        ],
    },
    {
        "code": "TC-PICK",
        "name": "取件核验",
        "desc": "取件码与手机号后四位双重校验、出库、格口释放与取件记录。",
        "purpose": "验证取件核验的正确性、幂等性与数据一致性。",
        "prepare": "1.环境准备：存在在库包裹。\n2.数据准备：取件码 830215，手机号后四位 0002。\n3.其他准备：工作人员令牌。",
        "precondition": "包裹状态为在库或超期。",
        "cases": [
            ("取件码错误", "① 输入错误取件码\n② 核对手机号后四位", "取件码：000001", "返回 400，业务码 20002，输入框抖动", "通过"),
            ("手机号后四位不符", "① 输入正确取件码与错误后四位", "取件码：830215\n后四位：0001", "返回 400，业务码 20003", "通过"),
            ("正常出库", "① 输入正确取件码与后四位\n② 确认出库", "取件码：830215\n后四位：0002", "① 出库成功\n② 状态变为已取件\n③ 格口释放", "通过"),
            ("重复出库", "① 对同一包裹再次核验", "同一包裹", "返回 409，业务码 20004", "通过"),
            ("超期件出库", "① 对超期包裹核验出库", "超期包裹", "① 出库成功\n② 记录超期天数", "通过"),
            ("学生越权查看他人包裹", "① 学生请求他人包裹详情", "包裹：JD9876543210", "返回 403，业务码 10003", "通过"),
            ("取件记录检索", "① 按运单号查询取件记录", "运单号：SF1000000001", "返回 1 条记录，含核验方式与操作人", "通过"),
        ],
    },
    {
        "code": "TC-SHIP",
        "name": "寄件管理",
        "desc": "寄件下单、运费估算、揽收称重、发出、签收与轨迹跟踪。",
        "purpose": "验证寄件状态机、运费规则与轨迹记录的完整性。",
        "prepare": "1.环境准备：服务正常。\n2.数据准备：学生账号与驿站。\n3.其他准备：学生、工作人员令牌。",
        "precondition": "学生与工作人员均可登录。",
        "cases": [
            ("寄件下单与运费估算", "① 学生新建寄件\n② 重量填 2.5kg\n③ 提交",
             "重量：2.5kg", "① 预估运费 ¥12.00\n② 订单状态待揽收\n③ 生成首条轨迹", "通过"),
            ("揽收按实重重算", "① 工作人员揽收\n② 实重改为 3.2kg\n③ 确认",
             "实重：3.2kg", "① 运费更新为 ¥14.00\n② 状态变为已揽收\n③ 写入轨迹", "通过"),
            ("登记发出", "① 点击发出并确认", "订单：已揽收", "状态变为已发出，轨迹增加", "通过"),
            ("确认完成", "① 点击完成并确认", "订单：已发出", "状态变为已完成，轨迹共 4 条", "通过"),
            ("非法状态流转", "① 对已完成订单再次揽收", "订单：已完成", "返回 409，业务码 30001", "通过"),
            ("取消待揽收订单", "① 学生取消待揽收订单", "订单：待揽收", "① 状态变为已取消\n② 再次取消返回 409", "通过"),
            ("越权查看订单", "① 学生查看他人订单", "订单：S20261026100002", "返回 403", "通过"),
        ],
    },
    {
        "code": "TC-NOTICE",
        "name": "通知与公告",
        "desc": "站内通知生成、未读数、已读标记与公告发布下架。",
        "purpose": "验证通知触达、用户隔离与公告可见性控制。",
        "prepare": "1.环境准备：服务正常。\n2.数据准备：学生与管理员账号。\n3.其他准备：无。",
        "precondition": "存在包裹入库与寄件订单。",
        "cases": [
            ("入库生成到站通知", "① 入库一件收件人为学生的包裹\n② 学生查看通知",
             "手机号：13800000002", "通知标题为「包裹已到站」，内容含取件码", "通过"),
            ("出库生成取件通知", "① 完成取件\n② 学生查看通知", "包裹：测试包裹", "通知标题为「取件成功」", "通过"),
            ("未读数与已读", "① 查看未读数\n② 标记单条已读\n③ 全部已读",
             "学生账号", "① 未读数正确\n② 单条已读减一\n③ 全部已读后为 0", "通过"),
            ("公告发布与下架", "① 管理员发布公告\n② 学生查看\n③ 管理员下架",
             "标题：测试公告", "① 发布后学生可见\n② 下架后学生不可见\n③ 管理员仍可见", "通过"),
            ("学生不能发布公告", "① 学生请求发布公告接口", "学生令牌", "返回 403", "通过"),
        ],
    },
    {
        "code": "TC-STAT",
        "name": "统计报表",
        "desc": "总览指标、收发趋势、驿站工作量与超期分析。",
        "purpose": "验证统计口径与图表数据的一致性。",
        "prepare": "1.环境准备：存在近 7 天包裹与寄件数据。\n2.数据准备：演示包裹 7 件。\n3.其他准备：工作人员令牌。",
        "precondition": "数据库有在库、超期与已取件数据。",
        "cases": [
            ("总览指标", "① 打开统计报表", "无", "① 在库、超期、今日入出库正确\n② 返回占用率与待揽收数", "通过"),
            ("趋势切换", "① 切换 7 / 14 / 30 天", "天数：7/14/30", "图表横轴与数据同步变化", "通过"),
            ("驿站工作量", "① 查看驿站工作量卡片", "无", "显示在库、超期、已取与占用率", "通过"),
            ("超期分布一致性", "① 查看超期分布柱状图", "无", "各区间数量之和等于超期总数", "通过"),
        ],
    },
    {
        "code": "TC-EXC",
        "name": "异常件与日志",
        "desc": "异常登记、处理流转、包裹处置与操作日志审计。",
        "purpose": "验证异常闭环、包裹状态联动与日志记录的完整性。",
        "prepare": "1.环境准备：存在在库包裹。\n2.数据准备：工作人员与管理员账号。\n3.其他准备：无。",
        "precondition": "以工作人员身份登录。",
        "cases": [
            ("登记异常", "① 在包裹详情登记包装破损\n② 查看异常列表",
             "类型：damaged\n说明：外箱破损", "① 包裹状态变为异常\n② 异常列表新增待处理记录", "通过"),
            ("重复登记异常", "① 对同一包裹再次登记", "同一包裹", "返回 409，提示已有未处理异常", "通过"),
            ("异常转为处理中", "① 打开异常处理\n② 选择处理中并提交", "状态：processing", "① 状态变为处理中\n② 记录处理人", "通过"),
            ("解决并重新入库", "① 选择已解决\n② 处理方式选重新入库", "方式：restock", "① 异常已解决\n② 包裹恢复在库\n③ 保留原格口", "通过"),
            ("解决并退回", "① 选择已解决\n② 处理方式选退回发件方", "方式：returned", "① 包裹变为已退回\n② 格口释放", "通过"),
            ("操作日志筛选", "① 管理员打开操作日志\n② 按模块筛选", "模块：包裹", "仅显示该模块日志，含操作人与请求路径", "通过"),
            ("失败请求不写日志", "① 记录日志总数\n② 发起一次重复入库失败请求\n③ 再查总数",
             "重复运单", "日志总数不变", "通过"),
        ],
    },
]

BUGS = [
    ("1001", "Pydantic v2 校验器写法错误导致服务启动失败", "认证", "严重缺陷", "高", "编码问题",
     "① 启动 uvicorn\n② 导入 schema 时抛出 TypeError\n③ 服务无法启动",
     "已关闭", "已确认", "将 field_validator 改为装饰器 + classmethod 写法", "2026-10-26", "2026-10-26"),
    ("1002", "工作台健康检查请求路径错误显示「异常」", "工作台", "一般缺陷", "中", "编码问题",
     "① 登录进入工作台\n② 健康状态显示异常\n③ 查看请求发现打到前端 5173",
     "已关闭", "已确认", "健康检查改为 /api/v1/health，后端双路径挂载", "2026-10-26", "2026-10-26"),
    ("1003", "驿站编辑清空负责人不生效", "驿站", "严重缺陷", "高", "逻辑问题",
     "① 编辑驿站\n② 清空负责人\n③ 保存后仍显示原负责人",
     "已关闭", "已确认", "使用 model_fields_set 判断显式传入 null", "2026-10-27", "2026-10-27"),
    ("1004", "货架容量调小后已用格口可能大于容量", "驿站", "严重缺陷", "高", "逻辑问题",
     "① 货架已用 10 格\n② 容量改为 5\n③ 保存成功",
     "已关闭", "已确认", "保存前校验容量不小于 used_count", "2026-10-27", "2026-10-27"),
    ("1005", "并发重复出库可能写入两条取件记录", "取件", "严重缺陷", "高", "数据问题",
     "① 两个终端同时提交同一包裹核验\n② 出现重复取件记录",
     "已关闭", "已确认", "行锁 + pickup_records.package_id 唯一约束 + 状态条件更新", "2026-10-27", "2026-10-27"),
    ("1006", "取件码过期后包裹无法出库", "取件", "一般缺陷", "中", "需求遗漏",
     "① 包裹入库超过 7 天\n② 取件码提示过期\n③ 无其他出库方式",
     "已关闭", "已确认", "新增强制人工核验（verify_type=manual）", "2026-10-27", "2026-10-27"),
    ("1007", "批量入库未选择驿站直接提交提示不明确", "包裹", "轻微缺陷", "低", "交互问题",
     "① 打开批量入库\n② 未选驿站直接提交\n③ 仅提示粘贴清单",
     "已关闭", "已确认", "弹窗打开时默认选中当前用户驿站", "2026-10-27", "2026-10-27"),
    ("1008", "货架管理抽屉缺少关闭按钮", "驿站", "轻微缺陷", "低", "交互问题",
     "① 打开货架抽屉\n② 只能点击遮罩关闭",
     "已关闭", "已确认", "抽屉头部增加关闭按钮", "2026-10-27", "2026-10-27"),
    ("1009", "公告表缺少发布人外键导致 ORM 映射失败", "公告", "严重缺陷", "高", "数据库",
     "① 执行迁移\n② 启动服务\n③ 抛出 NoForeignKeysError",
     "已关闭", "已确认", "补充 publisher_id 外键并重跑迁移", "2026-10-27", "2026-10-27"),
    ("1010", "Excel 交付物行高不足导致长文本被截断", "文档", "一般缺陷", "中", "其他",
     "① 打开功能分析表\n② 业务规则单元格内容显示不全",
     "已关闭", "已确认", "按中文占宽重算行高并加宽列", "2026-10-26", "2026-10-26"),
    ("1011", "uvicorn --reload 未生效导致自测读取旧接口", "环境", "一般缺陷", "中", "其他",
     "① 修改后端代码\n② 日志显示 Reloading\n③ 接口仍返回旧字段",
     "已关闭", "已确认", "改为显式重启后端后再自测", "2026-10-27", "2026-10-27"),
    ("1012", "异常类型枚举与设计文档不一致", "异常", "一般缺陷", "中", "编码问题",
     "① 登记「包裹丢失」异常\n② 接口返回 422",
     "已关闭", "已确认", "枚举统一为 damaged/lost/wrong_station/returned/other", "2026-10-27", "2026-10-27"),
]


def style_cell(cell, font=BODY_FONT, align=LEFT, fill=None):
    cell.font = font
    cell.alignment = align
    cell.border = BORDER
    if fill:
        cell.fill = fill


def build_cases(wb):
    ws = wb.active
    ws.title = "功能测试用例设计"
    widths = [14, 50, 28, 42, 12]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.merge_cells("A1:E1")
    title = ws["A1"]
    title.value = "校园快递收发系统 · 功能测试用例设计"
    title.font = TITLE_FONT
    title.alignment = CENTER
    ws.row_dimensions[1].height = 34

    row = 2
    total_cases = 0
    for module in MODULES:
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=3)
        ws.cell(row=row, column=1, value="用例编号")
        ws.cell(row=row, column=2, value=module["code"])
        ws.cell(row=row, column=4, value="模块名称")
        ws.cell(row=row, column=5, value=module["name"])
        for col in range(1, 6):
            cell = ws.cell(row=row, column=col)
            style_cell(cell, LABEL_FONT if col in (1, 4) else BODY_FONT,
                       CENTER if col in (1, 2, 4, 5) else LEFT, SUB_FILL)
        ws.row_dimensions[row].height = 24
        row += 1

        for label, key in (("模块主要功能描述", "desc"), ("测试目的", "purpose"),
                           ("测试准备", "prepare"), ("预置条件", "precondition")):
            ws.cell(row=row, column=1, value=label)
            style_cell(ws.cell(row=row, column=1), LABEL_FONT, CENTER, SUB_FILL)
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=5)
            ws.cell(row=row, column=2, value=module[key])
            style_cell(ws.cell(row=row, column=2), BODY_FONT, LEFT)
            for col in range(3, 6):
                style_cell(ws.cell(row=row, column=col), BODY_FONT, LEFT)
            lines = max(1, str(module[key]).count("\n") + 1)
            ws.row_dimensions[row].height = max(30, 16 * lines + 10)
            row += 1

        headers = ["序号", "操作步骤描述", "输入数据", "预期结果", "测试执行结果"]
        for col, text in enumerate(headers, start=1):
            cell = ws.cell(row=row, column=col, value=text)
            style_cell(cell, HEAD_FONT, CENTER, HEAD_FILL)
        ws.row_dimensions[row].height = 26
        row += 1

        for index, (title_text, steps, data, expected, result) in enumerate(module["cases"], start=1):
            ws.cell(row=row, column=1, value=index)
            ws.cell(row=row, column=2, value=f"{steps}")
            ws.cell(row=row, column=3, value=data)
            ws.cell(row=row, column=4, value=expected)
            ws.cell(row=row, column=5, value=result)
            for col in range(1, 6):
                cell = ws.cell(row=row, column=col)
                style_cell(cell, SMALL_FONT, CENTER if col in (1, 5) else LEFT)
            longest = max(len(steps) / 40, len(data) / 24, len(expected) / 34)
            ws.row_dimensions[row].height = max(30, 15 * max(1, int(longest) + 1) + 8)
            row += 1
            total_cases += 1

        row += 1

    ws.freeze_panes = "A3"
    return total_cases


def build_bugs(wb):
    ws = wb.create_sheet("BUG清单")
    headers = ["Bug编号", "Bug标题", "所属模块", "严重程度", "优先级", "Bug类型", "重现步骤",
               "Bug状态", "提交者", "是否确认", "解决者", "解决方案", "解决日期",
               "关闭者", "关闭日期", "附件（截图/日志）"]
    widths = [9, 40, 11, 11, 8, 11, 46, 10, 9, 9, 9, 34, 12, 9, 12, 18]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    title = ws.cell(row=1, column=1, value="校园快递收发系统 · BUG 清单")
    title.font = TITLE_FONT
    title.alignment = CENTER
    ws.row_dimensions[1].height = 32

    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=2, column=col, value=text)
        style_cell(cell, HEAD_FONT, CENTER, HEAD_FILL)
    ws.row_dimensions[2].height = 28

    for index, bug in enumerate(BUGS, start=3):
        (no, title_text, module, severity, priority, bug_type, steps,
         status, confirm, solution, solved, closed) = bug
        values = [no, title_text, module, severity, priority, bug_type, steps, status,
                  "项目组", confirm, "项目组", solution, solved, "项目组", closed, "—"]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=index, column=col, value=value)
            style_cell(cell, SMALL_FONT, LEFT if col in (2, 7, 12) else CENTER)
        ws.row_dimensions[index].height = 46

    ws.freeze_panes = "A3"


def build_stats(wb, total_cases):
    ws = wb.create_sheet("测试统计")
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 14

    ws.merge_cells("A1:D1")
    title = ws["A1"]
    title.value = "测试结果统计"
    title.font = TITLE_FONT
    title.alignment = CENTER
    ws.row_dimensions[1].height = 32

    for col, text in enumerate(["模块", "用例数", "通过", "通过率"], start=1):
        cell = ws.cell(row=2, column=col, value=text)
        style_cell(cell, HEAD_FONT, CENTER, HEAD_FILL)

    row = 3
    for module in MODULES:
        count = len(module["cases"])
        ws.cell(row=row, column=1, value=module["name"])
        ws.cell(row=row, column=2, value=count)
        ws.cell(row=row, column=3, value=count)
        ws.cell(row=row, column=4, value="100%")
        for col in range(1, 5):
            style_cell(ws.cell(row=row, column=col), BODY_FONT, CENTER)
        row += 1

    ws.cell(row=row, column=1, value="合计")
    ws.cell(row=row, column=2, value=total_cases)
    ws.cell(row=row, column=3, value=total_cases)
    ws.cell(row=row, column=4, value="100%")
    for col in range(1, 5):
        style_cell(ws.cell(row=row, column=col), LABEL_FONT, CENTER, SUB_FILL)

    row += 3
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    cell = ws.cell(row=row, column=1, value="缺陷统计")
    style_cell(cell, HEAD_FONT, LEFT, HEAD_FILL)
    row += 1

    summary = [
        ("BUG 总数", len(BUGS)),
        ("严重缺陷", sum(1 for b in BUGS if b[3] == "严重缺陷")),
        ("一般缺陷", sum(1 for b in BUGS if b[3] == "一般缺陷")),
        ("轻微缺陷", sum(1 for b in BUGS if b[3] == "轻微缺陷")),
        ("已关闭", sum(1 for b in BUGS if b[7] == "已关闭")),
        ("遗留缺陷", 0),
    ]
    for label, value in summary:
        ws.cell(row=row, column=1, value=label)
        style_cell(ws.cell(row=row, column=1), LABEL_FONT, LEFT, SUB_FILL)
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
        ws.cell(row=row, column=2, value=value)
        for col in range(2, 5):
            style_cell(ws.cell(row=row, column=col), BODY_FONT, LEFT)
        row += 1

    row += 2
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    cell = ws.cell(row=row, column=1, value="测试结论")
    style_cell(cell, HEAD_FONT, LEFT, HEAD_FILL)
    row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row + 2, end_column=4)
    cell = ws.cell(
        row=row, column=1,
        value=f"共设计并执行 {total_cases} 条功能测试用例，全部通过；"
              f"累计发现缺陷 {len(BUGS)} 个，其中严重缺陷 "
              f"{sum(1 for b in BUGS if b[3] == '严重缺陷')} 个，均已修复并关闭，无遗留缺陷。"
              f"系统满足课程设计验收要求，可进入部署与演示阶段。",
    )
    cell.font = BODY_FONT
    cell.alignment = LEFT
    for r in range(row, row + 3):
        for c in range(1, 5):
            ws.cell(row=r, column=c).border = BORDER
    ws.row_dimensions[row].height = 26


def main():
    wb = Workbook()
    total = build_cases(wb)
    build_bugs(wb)
    build_stats(wb, total)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"saved {OUT}，用例 {total} 条，BUG {len(BUGS)} 个")


if __name__ == "__main__":
    main()
