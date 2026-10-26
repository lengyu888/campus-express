# -*- coding: utf-8 -*-
"""将 docs/系统设计文档.md 转换为排版规范的 Word 文档。

运行：
    python scripts/build_docx.py
输出：
    docs/系统设计文档.docx
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parent.parent
MD = ROOT / "docs" / "系统设计文档.md"
OUT = ROOT / "docs" / "系统设计文档.docx"

FONT_SONG = "宋体"
FONT_HEI = "黑体"
FONT_ASCII = "Times New Roman"
FONT_CODE = "Consolas"
INK = RGBColor(0x1F, 0x29, 0x37)
BLUE = RGBColor(0x1F, 0x4E, 0x79)
GRAY = RGBColor(0x59, 0x59, 0x59)


# ---------------------------------------------------------------- 基础工具
def set_run_font(run, ascii_font=FONT_ASCII, ea_font=FONT_SONG, size=None,
                 bold=None, color=None):
    run.font.name = ascii_font
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), ascii_font)
    rfonts.set(qn("w:hAnsi"), ascii_font)
    rfonts.set(qn("w:eastAsia"), ea_font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def shade(element, color):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    element.append(shd)


def cell_shade(cell, color):
    shade(cell._tc.get_or_add_tcPr(), color)


def cell_text(cell, text, size=9, bold=False, color=INK, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(1)
    pf.space_after = Pt(1)
    pf.line_spacing = 1.0
    add_runs(p, text, size=size, bold=bold, color=color)


def set_cell_width(cell, cm):
    cell.width = Cm(cm)
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = tcPr.find(qn("w:tcW"))
    if tcW is None:
        tcW = OxmlElement("w:tcW")
        tcPr.append(tcW)
    tcW.set(qn("w:w"), str(int(cm * 567)))
    tcW.set(qn("w:type"), "dxa")


def set_table_borders(table, color="BFBFBF", sz=4):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tblPr.append(borders)


def set_table_cell_margins(table, left=0.08, right=0.08, top=0.03, bottom=0.03):
    tblPr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for k, v in (("left", left), ("right", right), ("top", top), ("bottom", bottom)):
        el = OxmlElement(f"w:{k}")
        el.set(qn("w:w"), str(int(v * 567)))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblPr.append(mar)


def fixed_layout(table):
    tblPr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def add_field(paragraph, instr, placeholder=""):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    run._r.append(fld_begin)

    run2 = paragraph.add_run()
    instr_el = OxmlElement("w:instrText")
    instr_el.set(qn("xml:space"), "preserve")
    instr_el.text = f" {instr} "
    run2._r.append(instr_el)

    run3 = paragraph.add_run()
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    run3._r.append(fld_sep)

    if placeholder:
        paragraph.add_run(placeholder)

    run4 = paragraph.add_run()
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run4._r.append(fld_end)
    return [run, run2, run3, run4]


def para_border_bottom(paragraph, color="1F4E79", sz=8):
    pPr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(sz))
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    pPr.append(borders)


def set_page_number_start(section, start=1):
    sectPr = section._sectPr
    pg = sectPr.find(qn("w:pgNumType"))
    if pg is None:
        pg = OxmlElement("w:pgNumType")
        sectPr.append(pg)
    pg.set(qn("w:start"), str(start))


# ---------------------------------------------------------------- 样式
def setup_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = FONT_ASCII
    normal.font.size = Pt(12)
    normal.font.color.rgb = INK
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_SONG)
    pf = normal.paragraph_format
    pf.line_spacing = 1.5
    pf.space_before = Pt(0)
    pf.space_after = Pt(3)
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT

    specs = {
        "Heading 1": (16, 18, 12),
        "Heading 2": (14, 12, 6),
        "Heading 3": (12, 10, 4),
    }
    for name, (size, before, after) in specs.items():
        st = doc.styles[name]
        st.font.name = FONT_ASCII
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = BLUE
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_HEI)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.line_spacing = 1.3
        st.paragraph_format.keep_with_next = True
        st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT


def setup_page(doc):
    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.54)
    sec.bottom_margin = Cm(2.54)
    sec.left_margin = Cm(3.0)
    sec.right_margin = Cm(2.5)
    sec.header_distance = Cm(1.5)
    sec.footer_distance = Cm(1.5)


def body_width_cm(doc):
    sec = doc.sections[-1]
    return (sec.page_width - sec.left_margin - sec.right_margin) / 360000


# ---------------------------------------------------------------- 富文本
TOKEN = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")


def add_runs(paragraph, text, size=12, bold=False, color=INK,
             ascii_font=FONT_ASCII, ea_font=FONT_SONG):
    for part in TOKEN.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, ascii_font, ea_font, size, True, color)
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, FONT_CODE, FONT_CODE, size - 1, bold, color)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, ascii_font, ea_font, size, bold, color)


def zh_quotes(text):
    out = []
    opening = True
    for ch in text:
        if ch == '"':
            out.append("“" if opening else "”")
            opening = not opening
        else:
            out.append(ch)
    return "".join(out)


# ---------------------------------------------------------------- Markdown 解析
def parse_markdown(path):
    lines = path.read_text(encoding="utf-8").splitlines()
    blocks = []
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if stripped.startswith("```"):
            lang = stripped[3:].strip()
            i += 1
            buf = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1
            blocks.append({"type": "code", "text": "\n".join(buf), "lang": lang})
            continue

        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            blocks.append({"type": "heading", "level": len(m.group(1)),
                           "text": zh_quotes(m.group(2))})
            i += 1
            continue

        m = re.match(r"^!\[(.*?)\]\((.*?)\)$", stripped)
        if m:
            blocks.append({"type": "image", "caption": m.group(1), "src": m.group(2)})
            i += 1
            continue

        if stripped.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                raw = lines[i].strip().strip("|")
                cells = [c.strip() for c in raw.split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    rows.append([zh_quotes(c) for c in cells])
                i += 1
            if rows:
                blocks.append({"type": "table", "rows": rows})
            continue

        m = re.match(r"^-\s+(.*)$", stripped)
        if m:
            blocks.append({"type": "bullet", "text": zh_quotes(m.group(1))})
            i += 1
            continue

        m = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if m:
            blocks.append({"type": "ordered", "num": m.group(1),
                           "text": zh_quotes(m.group(2))})
            i += 1
            continue

        blocks.append({"type": "para", "text": zh_quotes(stripped)})
        i += 1
    return blocks


# ---------------------------------------------------------------- 封面与前置页
def build_cover(doc):
    for _ in range(1):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("软件工程课程设计")
    set_run_font(run, FONT_ASCII, FONT_HEI, 26, True, BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(0)
    para_border_bottom(p, "1F4E79", 10)
    run = p.add_run(" ")
    set_run_font(run, FONT_ASCII, FONT_SONG, 2)

    for _ in range(2):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run("校园快递收发系统")
    set_run_font(run, FONT_ASCII, FONT_HEI, 34, True, INK)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("系统设计文档")
    set_run_font(run, FONT_ASCII, FONT_HEI, 24, False, BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    run = p.add_run("System Design Document")
    set_run_font(run, "Times New Roman", FONT_SONG, 13, False, GRAY)

    for _ in range(4):
        doc.add_paragraph()

    items = [
        ("专业班级", "____________________"),
        ("小组编号", "____________________"),
        ("小组成员", "____________________"),
        ("指导教师", "穆海伦"),
    ]
    for k, v in items:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(8)
        run = p.add_run(f"{k}：{v}")
        set_run_font(run, FONT_ASCII, FONT_SONG, 14, False, INK)

    for _ in range(2):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("2026 年 10 月")
    set_run_font(run, FONT_ASCII, FONT_SONG, 14, False, INK)


def build_front_matter(doc):
    doc.add_page_break()

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run("文档信息")
    set_run_font(run, FONT_ASCII, FONT_HEI, 16, True, BLUE)

    info = [
        ["文档名称", "校园快递收发系统系统设计文档", "文档编号", "CE-DESIGN-001"],
        ["版本", "V1.0", "密级", "内部"],
        ["编写", "项目组", "审核", "____________"],
        ["日期", "2026-10-26", "状态", "评审稿"],
    ]
    t = doc.add_table(rows=len(info), cols=4)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    fixed_layout(t)
    set_table_borders(t)
    set_table_cell_margins(t, 0.12, 0.12, 0.06, 0.06)
    widths = [2.6, 7.6, 2.6, 3.0]
    for c, w in enumerate(widths):
        t.columns[c].width = Cm(w)
    for r, row in enumerate(info):
        for c, val in enumerate(row):
            cell = t.cell(r, c)
            set_cell_width(cell, widths[c])
            bold = c % 2 == 0
            cell_text(cell, val, size=10, bold=bold,
                      align=WD_ALIGN_PARAGRAPH.CENTER if bold else WD_ALIGN_PARAGRAPH.LEFT)
            if bold:
                cell_shade(cell, "EAF2FA")

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run("修订记录")
    set_run_font(run, FONT_ASCII, FONT_HEI, 16, True, BLUE)

    rev = [
        ["版本", "日期", "修订人", "修订说明"],
        ["V1.0", "2026-10-26", "项目组", "初稿：总体结构、模块划分、接口设计、数据库设计、详细设计与部署设计"],
        ["", "", "", ""],
        ["", "", "", ""],
    ]
    t2 = doc.add_table(rows=len(rev), cols=4)
    t2.alignment = WD_TABLE_ALIGNMENT.CENTER
    t2.autofit = False
    fixed_layout(t2)
    set_table_borders(t2)
    set_table_cell_margins(t2, 0.1, 0.1, 0.06, 0.06)
    widths2 = [1.8, 3.0, 2.6, 8.4]
    for c, w in enumerate(widths2):
        t2.columns[c].width = Cm(w)
    for r, row in enumerate(rev):
        for c, val in enumerate(row):
            cell = t2.cell(r, c)
            set_cell_width(cell, widths2[c])
            header = r == 0
            cell_text(cell, val, size=9.5, bold=header,
                      align=WD_ALIGN_PARAGRAPH.CENTER if header or c < 3 else WD_ALIGN_PARAGRAPH.LEFT)
            if header:
                cell_shade(cell, "D9E2F3")
    repeat_header(t2.rows[0])


def build_toc(doc):
    doc.add_page_break()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run("目    录")
    set_run_font(run, FONT_ASCII, FONT_HEI, 18, True, BLUE)

    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    add_field(p, 'TOC \\o "1-3" \\h \\z \\u', "打开文档后，右键此处选择“更新域”生成目录")

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    run = p.add_run("提示：目录为 Word 域，首次打开如未显示页码，请全选后按 F9 或右键选择“更新域”。")
    set_run_font(run, FONT_SONG, FONT_SONG, 9, False, GRAY)


# ---------------------------------------------------------------- 页眉页脚
def build_header_footer(doc, section):
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False

    try:
        doc.styles["Header"].paragraph_format.tab_stops.clear_all()
    except KeyError:
        pass

    hp = section.header.paragraphs[0]
    hp.text = ""
    hp.paragraph_format.tab_stops.clear_all()
    hp.paragraph_format.tab_stops.add_tab_stop(Cm(15.4), 2)  # 2 = RIGHT
    run = hp.add_run("校园快递收发系统")
    set_run_font(run, FONT_ASCII, FONT_SONG, 9, False, GRAY)
    run = hp.add_run("\t系统设计文档")
    set_run_font(run, FONT_ASCII, FONT_SONG, 9, False, GRAY)
    para_border_bottom(hp, "BFBFBF", 4)

    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run("第 ")
    set_run_font(run, FONT_ASCII, FONT_SONG, 9, False, GRAY)
    add_field(fp, "PAGE", "1")
    run = fp.add_run(" 页 / 共 ")
    set_run_font(run, FONT_ASCII, FONT_SONG, 9, False, GRAY)
    add_field(fp, "NUMPAGES", "1")
    run = fp.add_run(" 页")
    set_run_font(run, FONT_ASCII, FONT_SONG, 9, False, GRAY)


# ---------------------------------------------------------------- 正文元素
def add_heading(doc, text, level, first_h1):
    if level == 1 and not first_h1:
        doc.add_page_break()
    p = doc.add_paragraph(style=f"Heading {level}")
    run = p.add_run(text)
    size = {1: 16, 2: 14, 3: 12}.get(level, 12)
    set_run_font(run, FONT_ASCII, FONT_HEI, size, True, BLUE)
    return p


def add_body_para(doc, text, indent=True):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(3)
    if indent:
        pf.first_line_indent = Pt(24)
    add_runs(p, text, size=12)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Cm(0.74)
    pf.first_line_indent = Cm(-0.37)
    pf.line_spacing = 1.5
    pf.space_after = Pt(2)
    add_runs(p, "• " + text, size=12)


def add_ordered(doc, num, text):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Cm(0.74)
    pf.first_line_indent = Cm(-0.37)
    pf.line_spacing = 1.5
    pf.space_after = Pt(2)
    add_runs(p, f"{num}. " + text, size=12)


def add_code(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.left_indent = Cm(0.4)
    pf.right_indent = Cm(0.2)
    pf.space_before = Pt(4)
    pf.space_after = Pt(6)
    pf.line_spacing = 1.0
    pf.keep_together = True
    pPr = p._p.get_or_add_pPr()
    shade(pPr, "F6F8FA")
    borders = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "4")
        el.set(qn("w:color"), "D0D7DE")
        borders.append(el)
    pPr.append(borders)

    lines = text.split("\n")
    for idx, line in enumerate(lines):
        run = p.add_run(line)
        set_run_font(run, FONT_CODE, FONT_CODE, 9, False, RGBColor(0x24, 0x29, 0x2E))
        if idx < len(lines) - 1:
            run.add_break()


def add_image(doc, src, caption):
    path = (MD.parent / src).resolve()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True

    with Image.open(path) as img:
        w_px, h_px = img.size
    max_w = 15.5
    max_h = 20.5
    width = max_w
    height = width * h_px / w_px
    if height > max_h:
        height = max_h
        width = height * w_px / h_px
    p.add_run().add_picture(str(path), width=Cm(width))

    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_before = Pt(0)
    cp.paragraph_format.space_after = Pt(10)
    run = cp.add_run(caption)
    set_run_font(run, FONT_ASCII, FONT_SONG, 10, True, INK)


def add_table(doc, rows, caption):
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_before = Pt(6)
    cp.paragraph_format.space_after = Pt(3)
    cp.paragraph_format.keep_with_next = True
    run = cp.add_run(caption)
    set_run_font(run, FONT_ASCII, FONT_SONG, 10, True, INK)

    n_cols = max(len(r) for r in rows)
    rows = [r + [""] * (n_cols - len(r)) for r in rows]
    table = doc.add_table(rows=len(rows), cols=n_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    fixed_layout(table)
    set_table_borders(table)
    set_table_cell_margins(table)

    def disp_len(s):
        return sum(2 if ord(ch) > 127 else 1 for ch in s)

    weights = []
    for c in range(n_cols):
        mx = max(disp_len(r[c]) for r in rows)
        weights.append(max(4, min(mx, 34)))
    total = sum(weights)
    avail = 15.8
    widths = [max(1.6, avail * w / total) for w in weights]
    scale = avail / sum(widths)
    widths = [w * scale for w in widths]
    for c, w in enumerate(widths):
        table.columns[c].width = Cm(w)

    for r_idx, row in enumerate(rows):
        tr = table.rows[r_idx]
        cant_split(tr)
        for c_idx, val in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            set_cell_width(cell, widths[c_idx])
            header = r_idx == 0
            cell_text(cell, val, size=8.5, bold=header,
                      align=WD_ALIGN_PARAGRAPH.CENTER if header else WD_ALIGN_PARAGRAPH.LEFT)
            if header:
                cell_shade(cell, "D9E2F3")
            cell.vertical_alignment = 1
    repeat_header(table.rows[0])
    if len(rows) <= 10:
        for row in table.rows[:-1]:
            for cell in row.cells:
                for para in cell.paragraphs:
                    para.paragraph_format.keep_with_next = True

    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(4)
    sp.paragraph_format.line_spacing = 0.8


def add_table_caption_number(text, chapter, seq):
    clean = re.sub(r"^[\d.]+\s*", "", text).strip()
    clean = re.sub(r"^附录\s*[A-Z]\s*", "", clean).strip()
    return f"表 {chapter}-{seq} {clean}"


# ---------------------------------------------------------------- 主流程
def main():
    blocks = parse_markdown(MD)
    doc = Document()
    setup_page(doc)
    setup_styles(doc)
    doc.core_properties.title = "校园快递收发系统系统设计文档"
    doc.core_properties.author = "项目组"
    doc.core_properties.subject = "软件工程课程设计"

    build_cover(doc)
    build_front_matter(doc)
    build_toc(doc)

    body = doc.add_section(WD_SECTION.NEW_PAGE)
    body.top_margin = Cm(2.54)
    body.bottom_margin = Cm(2.54)
    body.left_margin = Cm(3.0)
    body.right_margin = Cm(2.5)
    set_page_number_start(body, 1)
    build_header_footer(doc, body)

    chapter = "0"
    table_seq = 0
    current_heading = ""
    first_h1 = True

    for idx, b in enumerate(blocks):
        t = b["type"]
        if t == "heading":
            if b["level"] == 1:
                m = re.match(r"^(\d+)", b["text"])
                ma = re.match(r"^附录\s*([A-Z])", b["text"])
                if ma:
                    chapter = ma.group(1)
                elif m:
                    chapter = m.group(1)
                table_seq = 0
            current_heading = b["text"]
            add_heading(doc, b["text"], b["level"], first_h1)
            if b["level"] == 1:
                first_h1 = False
        elif t == "para":
            if re.match(r"^(GET|POST|PUT|PATCH|DELETE)\s+/", b["text"]):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.left_indent = Cm(0.4)
                add_runs(p, b["text"], size=10.5, bold=True, color=BLUE,
                         ascii_font=FONT_CODE, ea_font=FONT_CODE)
            else:
                para = add_body_para(doc, b["text"])
                nxt = blocks[idx + 1] if idx + 1 < len(blocks) else None
                if nxt and nxt["type"] == "image":
                    para.paragraph_format.keep_with_next = True
        elif t == "bullet":
            add_bullet(doc, b["text"])
        elif t == "ordered":
            add_ordered(doc, b["num"], b["text"])
        elif t == "code":
            add_code(doc, b["text"])
        elif t == "image":
            add_image(doc, b["src"], b["caption"])
        elif t == "table":
            table_seq += 1
            add_table(doc, b["rows"], add_table_caption_number(current_heading, chapter, table_seq))

    doc.save(OUT)
    print("saved", OUT)


if __name__ == "__main__":
    main()
