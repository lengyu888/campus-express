# -*- coding: utf-8 -*-
"""生成业务流程图 Word 文档（需求分析阶段交付物）。

输出：docs/交付物/业务流程图.docx
"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parent.parent
ASSETS = ROOT / "docs" / "assets"
OUT = ROOT / "docs" / "交付物" / "业务流程图.docx"

BLUE = RGBColor(0x1F, 0x4E, 0x79)
INK = RGBColor(0x1F, 0x29, 0x37)


def set_font(run, ea="宋体", ascii_font="Times New Roman", size=12, bold=False, color=INK):
    run.font.name = ascii_font
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        from docx.oxml import OxmlElement
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), ascii_font)
    rfonts.set(qn("w:hAnsi"), ascii_font)
    rfonts.set(qn("w:eastAsia"), ea)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def main():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.2)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.8)
    sec.right_margin = Cm(2.6)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.paragraph_format.line_spacing = 1.4
    normal.paragraph_format.space_after = Pt(2)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("校园快递收发系统 · 业务流程图")
    set_font(r, "黑体", size=20, bold=True, color=BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("软件工程课程设计 · 需求分析阶段")
    set_font(r, size=9.5, color=RGBColor(0x59, 0x59, 0x59))

    items = [
        ("一、收件业务流程", "图 1 收件业务流程", "fig_flow_business_recv.png",
         "快递到站后由驿站工作人员登记运单，系统校验并生成取件码、推送通知；收件人到站出示取件码与手机号后四位，"
         "核验通过后完成出库并释放格口。"),
        ("二、寄件业务流程", "图 2 寄件业务流程", "fig_flow_business_send.png",
         "学生在线填写寄件信息并下单，驿站揽收称重、确认运费并登记轨迹，系统负责状态流转与通知推送。"),
    ]

    for idx, (title, caption, img, desc) in enumerate(items):
        p = doc.add_paragraph()
        if idx > 0:
            p.paragraph_format.page_break_before = True
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(title)
        set_font(r, "黑体", size=13, bold=True, color=BLUE)

        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Pt(21)
        p.paragraph_format.line_spacing = 1.3
        r = p.add_run(desc)
        set_font(r, size=10)

        path = ASSETS / img
        with Image.open(path) as im:
            w, h = im.size
        width = 14.2
        height = width * h / w
        if height > 20.2:
            height = 20.2
            width = height * w / h

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(3)
        p.add_run().add_picture(str(path), width=Cm(width))

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(caption)
        set_font(r, size=10, bold=True)

    doc.core_properties.title = "校园快递收发系统业务流程图"
    doc.core_properties.author = "项目组"
    doc.save(OUT)
    print("saved", OUT)


if __name__ == "__main__":
    main()
