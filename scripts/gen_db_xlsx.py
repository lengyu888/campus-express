# -*- coding: utf-8 -*-
"""按《数据库表设计参考》格式生成数据库表结构 Excel。

输出：docs/交付物/数据库表结构.xlsx
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "docs" / "交付物"
OUT.mkdir(parents=True, exist_ok=True)

FONT = "宋体"
THIN = Side(style="thin", color="808080")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="D9E2F3")
TITLE_FILL = PatternFill("solid", fgColor="1F4E79")

HEADERS = ["字段名称", "字段类型", "空", "默认值", "约束", "主/外键", "索引", "描述信息", "数据来源", "取值规则", "备注"]
WIDTHS = [20, 18, 6, 20, 18, 9, 7, 28, 16, 28, 14]

# (字段, 类型, 空, 默认值, 约束, 主/外键, 索引, 描述, 数据来源, 取值规则)
TABLES = [
    ("users", "用户表", [
        ("id", "BIGINT UNSIGNED", "", "AUTO_INCREMENT", "NOT NULL", "PK", "√", "用户主键", "系统生成", "自增"),
        ("student_no", "VARCHAR(20)", "√", "NULL", "UNIQUE", "", "√", "学号或工号", "用户输入", "管理员可为空"),
        ("name", "VARCHAR(50)", "", "", "NOT NULL", "", "", "真实姓名", "用户输入", "2—20 个字符"),
        ("phone", "VARCHAR(11)", "", "", "NOT NULL, UNIQUE", "", "√", "手机号", "用户输入", "11 位数字"),
        ("password_hash", "VARCHAR(100)", "", "", "NOT NULL", "", "", "密码哈希", "系统生成", "bcrypt，禁止明文"),
        ("role", "VARCHAR(16)", "", "'student'", "NOT NULL", "", "√", "角色", "管理员设置", "student/staff/admin"),
        ("college", "VARCHAR(100)", "√", "NULL", "", "", "", "学院或部门", "用户输入", ""),
        ("avatar_url", "VARCHAR(255)", "√", "NULL", "", "", "", "头像地址", "用户上传", "相对路径或 URL"),
        ("status", "TINYINT", "", "1", "NOT NULL", "", "√", "账号状态", "管理员设置", "1 正常；0 禁用"),
        ("last_login_at", "DATETIME", "√", "NULL", "", "", "", "最近登录时间", "系统生成", ""),
        ("created_at", "DATETIME", "", "CURRENT_TIMESTAMP", "NOT NULL", "", "", "创建时间", "系统生成", ""),
        ("updated_at", "DATETIME", "", "CURRENT_TIMESTAMP", "NOT NULL", "", "", "更新时间", "系统生成", "ON UPDATE"),
    ]),
    ("stations", "驿站表", [
        ("id", "BIGINT UNSIGNED", "", "AUTO_INCREMENT", "NOT NULL", "PK", "√", "驿站主键", "系统生成", "自增"),
        ("code", "VARCHAR(20)", "", "", "NOT NULL, UNIQUE", "", "√", "驿站编码", "管理员录入", "ST- 开头"),
        ("name", "VARCHAR(50)", "", "", "NOT NULL", "", "", "驿站名称", "管理员录入", ""),
        ("address", "VARCHAR(200)", "", "", "NOT NULL", "", "", "详细地址", "管理员录入", ""),
        ("phone", "VARCHAR(11)", "√", "NULL", "", "", "", "联系电话", "管理员录入", "11 位数字"),
        ("open_time", "TIME", "√", "NULL", "", "", "", "营业开始时间", "管理员录入", "HH:mm:ss"),
        ("close_time", "TIME", "√", "NULL", "", "", "", "营业结束时间", "管理员录入", "晚于 open_time"),
        ("capacity", "INT", "", "0", "NOT NULL", "", "", "总格口数", "系统汇总", "≥0"),
        ("manager_id", "BIGINT UNSIGNED", "√", "NULL", "", "FK", "", "驿站负责人", "管理员指定", "引用 users.id"),
        ("status", "TINYINT", "", "1", "NOT NULL", "", "√", "驿站状态", "管理员设置", "1 营业；0 停用"),
        ("created_at", "DATETIME", "", "CURRENT_TIMESTAMP", "NOT NULL", "", "", "创建时间", "系统生成", ""),
        ("updated_at", "DATETIME", "", "CURRENT_TIMESTAMP", "NOT NULL", "", "", "更新时间", "系统生成", "ON UPDATE"),
    ]),
    ("shelves", "货架表", [
        ("id", "BIGINT UNSIGNED", "", "AUTO_INCREMENT", "NOT NULL", "PK", "√", "货架主键", "系统生成", "自增"),
        ("station_id", "BIGINT UNSIGNED", "", "", "NOT NULL", "FK", "√", "所属驿站", "管理员选择", "引用 stations.id"),
        ("code", "VARCHAR(20)", "", "", "NOT NULL", "", "√", "货架编号", "管理员录入", "驿站内唯一，如 A-03"),
        ("zone", "VARCHAR(20)", "√", "NULL", "", "", "", "分区", "管理员录入", "A—Z"),
        ("capacity", "INT", "", "20", "NOT NULL", "", "", "格口数量", "管理员录入", "1—200"),
        ("used_count", "INT", "", "0", "NOT NULL", "", "", "已用格口数", "系统维护", "0≤used≤capacity"),
        ("status", "TINYINT", "", "1", "NOT NULL", "", "", "货架状态", "管理员设置", "1 可用；0 停用"),
        ("created_at", "DATETIME", "", "CURRENT_TIMESTAMP", "NOT NULL", "", "", "创建时间", "系统生成", ""),
        ("updated_at", "DATETIME", "", "CURRENT_TIMESTAMP", "NOT NULL", "", "", "更新时间", "系统生成", "ON UPDATE"),
    ]),
    ("packages", "快递包裹表", [
        ("id", "BIGINT UNSIGNED", "", "AUTO_INCREMENT", "NOT NULL", "PK", "√", "包裹主键", "系统生成", "自增"),
        ("tracking_no", "VARCHAR(40)", "", "", "NOT NULL, UNIQUE", "", "√", "运单号", "扫码/录入", "全局唯一"),
        ("station_id", "BIGINT UNSIGNED", "", "", "NOT NULL", "FK", "√", "所属驿站", "工作人员选择", "引用 stations.id"),
        ("shelf_id", "BIGINT UNSIGNED", "√", "NULL", "", "FK", "√", "存放货架", "系统推荐", "出库后置空"),
        ("recipient_id", "BIGINT UNSIGNED", "√", "NULL", "", "FK", "√", "收件人", "系统匹配", "未匹配时为空"),
        ("recipient_name", "VARCHAR(50)", "√", "NULL", "", "", "", "收件人姓名", "运单信息", ""),
        ("recipient_phone", "VARCHAR(11)", "", "", "NOT NULL", "", "√", "收件人手机号", "扫码/录入", "11 位数字"),
        ("courier_company", "VARCHAR(30)", "", "", "NOT NULL", "", "", "快递公司", "工作人员选择", "参见数据列表"),
        ("package_type", "VARCHAR(16)", "", "'normal'", "NOT NULL", "", "", "包裹类型", "工作人员选择", "normal/fragile/large/doc"),
        ("weight", "DECIMAL(6,2)", "√", "NULL", "", "", "", "重量（kg）", "工作人员录入", "0.01—100.00"),
        ("status", "VARCHAR(16)", "", "'in_stock'", "NOT NULL", "", "√", "包裹状态", "系统维护", "in_stock/picked/overdue/returned/exception"),
        ("pickup_code", "VARCHAR(6)", "√", "NULL", "", "", "√", "取件码", "系统生成", "6 位数字，站内唯一"),
        ("code_expire_at", "DATETIME", "√", "NULL", "", "", "√", "取件码失效时间", "系统生成", "入库 +7 天"),
        ("stored_at", "DATETIME", "", "", "NOT NULL", "", "", "入库时间", "系统生成", "超期判定基准"),
        ("picked_at", "DATETIME", "√", "NULL", "", "", "", "取件时间", "系统生成", ""),
        ("overdue_days", "INT", "", "0", "NOT NULL", "", "", "超期天数", "系统计算", "≥0"),
        ("remark", "VARCHAR(255)", "√", "NULL", "", "", "", "备注", "工作人员录入", ""),
        ("created_by", "BIGINT UNSIGNED", "", "", "NOT NULL", "FK", "", "入库操作员", "当前登录用户", "引用 users.id"),
        ("created_at", "DATETIME", "", "CURRENT_TIMESTAMP", "NOT NULL", "", "", "创建时间", "系统生成", ""),
        ("updated_at", "DATETIME", "", "CURRENT_TIMESTAMP", "NOT NULL", "", "", "更新时间", "系统生成", "ON UPDATE"),
    ]),
    ("pickup_records", "取件记录表", [
        ("id", "BIGINT UNSIGNED", "", "AUTO_INCREMENT", "NOT NULL", "PK", "√", "记录主键", "系统生成", "自增"),
        ("package_id", "BIGINT UNSIGNED", "", "", "NOT NULL, UNIQUE", "FK", "√", "所属包裹", "系统带入", "一件包裹仅一条"),
        ("operator_id", "BIGINT UNSIGNED", "", "", "NOT NULL", "FK", "√", "出库操作员", "当前登录用户", "引用 users.id"),
        ("pickup_code", "VARCHAR(6)", "", "", "NOT NULL", "", "", "取件码快照", "系统带入", "6 位数字"),
        ("phone_suffix", "CHAR(4)", "", "", "NOT NULL", "", "", "手机号后四位", "工作人员录入", "4 位数字"),
        ("verify_type", "VARCHAR(16)", "", "'code_phone'", "NOT NULL", "", "", "核验方式", "工作人员选择", "code_phone/manual"),
        ("picked_at", "DATETIME", "", "", "NOT NULL", "", "√", "取件时间", "系统生成", ""),
        ("remark", "VARCHAR(200)", "√", "NULL", "", "", "", "备注", "工作人员录入", ""),
        ("created_at", "DATETIME", "", "CURRENT_TIMESTAMP", "NOT NULL", "", "", "创建时间", "系统生成", ""),
    ]),
    ("shipment_orders", "寄件订单表", [
        ("id", "BIGINT UNSIGNED", "", "AUTO_INCREMENT", "NOT NULL", "PK", "√", "订单主键", "系统生成", "自增"),
        ("order_no", "VARCHAR(32)", "", "", "NOT NULL, UNIQUE", "", "√", "订单号", "系统生成", "S+yyyyMMdd+6 位序号"),
        ("sender_id", "BIGINT UNSIGNED", "", "", "NOT NULL", "FK", "√", "寄件人", "当前登录用户", "引用 users.id"),
        ("station_id", "BIGINT UNSIGNED", "", "", "NOT NULL", "FK", "√", "受理驿站", "学生选择", "引用 stations.id"),
        ("receiver_name", "VARCHAR(50)", "", "", "NOT NULL", "", "", "收件人姓名", "学生录入", ""),
        ("receiver_phone", "VARCHAR(11)", "", "", "NOT NULL", "", "", "收件人电话", "学生录入", "11 位数字"),
        ("receiver_address", "VARCHAR(200)", "", "", "NOT NULL", "", "", "收件地址", "学生录入", ""),
        ("courier_company", "VARCHAR(30)", "√", "NULL", "", "", "", "快递公司", "工作人员选择", "揽收时指定"),
        ("item_name", "VARCHAR(100)", "", "", "NOT NULL", "", "", "物品名称", "学生录入", ""),
        ("weight", "DECIMAL(6,2)", "√", "NULL", "", "", "", "实际重量", "工作人员录入", "称重后确认"),
        ("freight", "DECIMAL(8,2)", "√", "NULL", "", "", "", "运费（元）", "系统计算", "首重 8 元，续重 2 元/kg"),
        ("status", "VARCHAR(16)", "", "'pending'", "NOT NULL", "", "√", "订单状态", "系统维护", "pending/accepted/shipped/completed/cancelled"),
        ("remark", "VARCHAR(255)", "√", "NULL", "", "", "", "备注", "学生录入", ""),
        ("created_at", "DATETIME", "", "CURRENT_TIMESTAMP", "NOT NULL", "", "", "下单时间", "系统生成", ""),
        ("updated_at", "DATETIME", "", "CURRENT_TIMESTAMP", "NOT NULL", "", "", "更新时间", "系统生成", "ON UPDATE"),
    ]),
    ("shipment_tracks", "寄件轨迹表", [
        ("id", "BIGINT UNSIGNED", "", "AUTO_INCREMENT", "NOT NULL", "PK", "√", "轨迹主键", "系统生成", "自增"),
        ("order_id", "BIGINT UNSIGNED", "", "", "NOT NULL", "FK", "√", "所属订单", "系统带入", "引用 shipment_orders.id"),
        ("node", "VARCHAR(50)", "", "", "NOT NULL", "", "", "节点名称", "工作人员录入", "如驿站揽收"),
        ("status", "VARCHAR(16)", "", "", "NOT NULL", "", "", "节点状态", "系统带入", "对应订单状态"),
        ("description", "VARCHAR(200)", "√", "NULL", "", "", "", "节点描述", "工作人员录入", ""),
        ("operator_id", "BIGINT UNSIGNED", "√", "NULL", "", "FK", "", "操作人", "当前登录用户", "引用 users.id"),
        ("occurred_at", "DATETIME", "", "", "NOT NULL", "", "√", "发生时间", "工作人员录入", ""),
        ("created_at", "DATETIME", "", "CURRENT_TIMESTAMP", "NOT NULL", "", "", "创建时间", "系统生成", ""),
    ]),
    ("notifications", "通知表", [
        ("id", "BIGINT UNSIGNED", "", "AUTO_INCREMENT", "NOT NULL", "PK", "√", "通知主键", "系统生成", "自增"),
        ("user_id", "BIGINT UNSIGNED", "", "", "NOT NULL", "FK", "√", "接收人", "系统带入", "引用 users.id"),
        ("type", "VARCHAR(16)", "", "", "NOT NULL", "", "√", "通知类型", "系统生成", "pickup/overdue/shipment/system"),
        ("title", "VARCHAR(100)", "", "", "NOT NULL", "", "", "通知标题", "系统生成", ""),
        ("content", "VARCHAR(500)", "", "", "NOT NULL", "", "", "通知内容", "系统生成", ""),
        ("package_id", "BIGINT UNSIGNED", "√", "NULL", "", "", "", "关联包裹", "系统带入", "可为空"),
        ("order_id", "BIGINT UNSIGNED", "√", "NULL", "", "", "", "关联订单", "系统带入", "可为空"),
        ("channel", "VARCHAR(16)", "", "'inapp'", "NOT NULL", "", "", "发送渠道", "系统设置", "inapp/sms"),
        ("is_read", "TINYINT", "", "0", "NOT NULL", "", "√", "是否已读", "用户操作", "0 未读；1 已读"),
        ("sent_at", "DATETIME", "√", "NULL", "", "", "", "发送时间", "系统生成", ""),
        ("created_at", "DATETIME", "", "CURRENT_TIMESTAMP", "NOT NULL", "", "", "创建时间", "系统生成", ""),
    ]),
    ("announcements", "公告表", [
        ("id", "BIGINT UNSIGNED", "", "AUTO_INCREMENT", "NOT NULL", "PK", "√", "公告主键", "系统生成", "自增"),
        ("title", "VARCHAR(100)", "", "", "NOT NULL", "", "", "公告标题", "管理员录入", ""),
        ("content", "TEXT", "", "", "NOT NULL", "", "", "公告正文", "管理员录入", ""),
        ("publisher_id", "BIGINT UNSIGNED", "", "", "NOT NULL", "FK", "", "发布人", "当前登录用户", "引用 users.id"),
        ("status", "TINYINT", "", "1", "NOT NULL", "", "√", "公告状态", "管理员设置", "1 已发布；0 草稿/下架"),
        ("publish_at", "DATETIME", "√", "NULL", "", "", "√", "发布时间", "系统生成", ""),
        ("created_at", "DATETIME", "", "CURRENT_TIMESTAMP", "NOT NULL", "", "", "创建时间", "系统生成", ""),
        ("updated_at", "DATETIME", "", "CURRENT_TIMESTAMP", "NOT NULL", "", "", "更新时间", "系统生成", "ON UPDATE"),
    ]),
    ("exception_records", "异常记录表", [
        ("id", "BIGINT UNSIGNED", "", "AUTO_INCREMENT", "NOT NULL", "PK", "√", "异常主键", "系统生成", "自增"),
        ("package_id", "BIGINT UNSIGNED", "", "", "NOT NULL", "FK", "√", "关联包裹", "系统带入", "引用 packages.id"),
        ("type", "VARCHAR(16)", "", "", "NOT NULL", "", "√", "异常类型", "工作人员选择", "damaged/lost/wrong_station/returned/other"),
        ("description", "VARCHAR(300)", "", "", "NOT NULL", "", "", "异常描述", "工作人员录入", ""),
        ("handler_id", "BIGINT UNSIGNED", "√", "NULL", "", "FK", "", "处理人", "管理员指定", "引用 users.id"),
        ("status", "VARCHAR(16)", "", "'open'", "NOT NULL", "", "√", "处理状态", "系统维护", "open/processing/resolved"),
        ("resolved_at", "DATETIME", "√", "NULL", "", "", "", "完成时间", "系统生成", ""),
        ("created_at", "DATETIME", "", "CURRENT_TIMESTAMP", "NOT NULL", "", "", "创建时间", "系统生成", ""),
        ("updated_at", "DATETIME", "", "CURRENT_TIMESTAMP", "NOT NULL", "", "", "更新时间", "系统生成", "ON UPDATE"),
    ]),
    ("operation_logs", "操作日志表", [
        ("id", "BIGINT UNSIGNED", "", "AUTO_INCREMENT", "NOT NULL", "PK", "√", "日志主键", "系统生成", "自增"),
        ("user_id", "BIGINT UNSIGNED", "√", "NULL", "", "", "√", "操作人", "当前登录用户", "系统任务为空"),
        ("module", "VARCHAR(30)", "", "", "NOT NULL", "", "√", "模块标识", "系统生成", "如 package"),
        ("action", "VARCHAR(50)", "", "", "NOT NULL", "", "", "操作动作", "系统生成", "如 inbound/pickup"),
        ("target_type", "VARCHAR(30)", "√", "NULL", "", "", "", "目标类型", "系统生成", ""),
        ("target_id", "BIGINT UNSIGNED", "√", "NULL", "", "", "", "目标主键", "系统生成", ""),
        ("detail", "JSON", "√", "NULL", "", "", "", "参数快照", "系统生成", "JSON 格式"),
        ("ip", "VARCHAR(45)", "√", "NULL", "", "", "", "客户端 IP", "系统生成", "兼容 IPv6"),
        ("created_at", "DATETIME", "", "CURRENT_TIMESTAMP", "NOT NULL", "", "√", "操作时间", "系统生成", ""),
    ]),
]


def row_height(vals):
    lines = 1
    for i, v in enumerate(vals):
        cpl = max(WIDTHS[i] / 1.9, 4)
        eff = sum(1 if ord(ch) > 127 else 0.55 for ch in str(v))
        lines = max(lines, int(eff / cpl) + 1)
    return max(20, lines * 13 + 5)


def style(c, bold=False, fill=None, align="left", size=10):
    c.font = Font(name=FONT, size=size, bold=bold)
    c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
    c.border = BORDER
    if fill:
        c.fill = fill


def main():
    wb = Workbook()
    cover = wb.active
    cover.title = "表清单"
    cover.merge_cells("A1:D1")
    t = cover.cell(row=1, column=1, value="校园快递收发系统 数据库表清单")
    t.font = Font(name=FONT, size=14, bold=True, color="FFFFFF")
    t.fill = TITLE_FILL
    t.alignment = Alignment(horizontal="center", vertical="center")
    cover.row_dimensions[1].height = 30

    head = ["序号", "表名", "中文名", "说明"]
    for i, h in enumerate(head, 1):
        style(cover.cell(row=2, column=i, value=h), bold=True, fill=HEAD_FILL, align="center")
    notes = {
        "users": "三类角色统一账号，含登录与状态",
        "stations": "多驿站模型，支持容量与营业时间",
        "shelves": "驿站下的货架与格口占用",
        "packages": "运单、取件码与包裹状态",
        "pickup_records": "出库核验记录，包裹唯一",
        "shipment_orders": "寄件订单与状态流转",
        "shipment_tracks": "寄件轨迹节点",
        "notifications": "取件、超期、寄件通知",
        "announcements": "管理员发布的公告",
        "exception_records": "异常件登记与处理",
        "operation_logs": "关键操作审计日志",
    }
    for i, (name, cn, _) in enumerate(TABLES, 3):
        style(cover.cell(row=i, column=1, value=i - 2), align="center")
        style(cover.cell(row=i, column=2, value=name))
        style(cover.cell(row=i, column=3, value=cn))
        style(cover.cell(row=i, column=4, value=notes.get(name, "")))
        cover.row_dimensions[i].height = 22
    cover.column_dimensions["A"].width = 8
    cover.column_dimensions["B"].width = 26
    cover.column_dimensions["C"].width = 24
    cover.column_dimensions["D"].width = 60

    ws = wb.create_sheet("数据库表结构")
    for i, w in enumerate(WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    row = 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=11)
    c = ws.cell(row=row, column=1, value="【校园快递收发系统】数据库表结构")
    c.font = Font(name=FONT, size=14, bold=True, color="FFFFFF")
    c.fill = TITLE_FILL
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[row].height = 30
    row += 2

    for name, cn, fields in TABLES:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=11)
        c = ws.cell(row=row, column=1, value=f"[{name}]－{cn}")
        c.font = Font(name=FONT, size=11, bold=True)
        c.fill = PatternFill("solid", fgColor="EAF2FA")
        c.alignment = Alignment(horizontal="left", vertical="center")
        c.border = BORDER
        ws.row_dimensions[row].height = 24
        row += 1

        for i, h in enumerate(HEADERS, 1):
            style(ws.cell(row=row, column=i, value=h), bold=True, fill=HEAD_FILL, align="center")
        ws.row_dimensions[row].height = 22
        row += 1

        for f in fields:
            vals = list(f) + [""]
            vals = vals[:11]
            for i, v in enumerate(vals, 1):
                style(ws.cell(row=row, column=i, value=v),
                      align="center" if i in (3, 6, 7) else "left")
            ws.row_dimensions[row].height = row_height(vals)
            row += 1
        row += 1

    ws.freeze_panes = "A4"
    path = OUT / "数据库表结构.xlsx"
    wb.save(path)
    print("saved", path)


if __name__ == "__main__":
    main()
